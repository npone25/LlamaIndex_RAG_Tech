# Tạo dữ liệu mẫu cho thí nghiệm 4.1.1 (chạy 1 lần)
# Cài đặt: pip install python-docx
#
# Sau khi chạy, thư mục thi_nghiem/Data_for_exp có dạng:
#   Data_for_exp/
#     bao_hanh/      quy_dinh_bao_hanh.md   (file mới tạo)
#     doi_tra/       chinh_sach_doi_tra.docx (file mới tạo)
#     tip_su_dung/   so_sanh_nhanh.csv, thong_so_mau.json, gioi_thieu.html (file mới tạo)
#
# Hai file có sẵn của bạn tự chép vào:
#   iphone-18-pro-max-info.pdf          -> Data_for_exp/bao_hanh/
#   ch4_exp_data_SimpleDirector.txt     -> Data_for_exp/tip_su_dung/

import json
import sys
from pathlib import Path

from docx import Document as WordDocument

sys.stdout.reconfigure(encoding="utf-8")  # để in được tiếng Việt trên Windows

root = Path(__file__).resolve().parent.parent / "thi_nghiem" / "Data_for_exp"

# ----------------------------------------------------------------------
# Tạo 3 thư mục con (tên thư mục sẽ trở thành doc_type ở bước 4)
# ----------------------------------------------------------------------
for folder_name in ["bao_hanh", "doi_tra", "tip_su_dung"]:
    (root / folder_name).mkdir(parents=True, exist_ok=True)

# ----------------------------------------------------------------------
# 1. File .md (có 2 tiêu đề)
# ----------------------------------------------------------------------
md_text = (
    "# Quy định bảo hành\n"
    "Thời gian bảo hành tiêu chuẩn là 12 tháng kể từ ngày kích hoạt.\n"
    "\n"
    "# Điều kiện loại trừ\n"
    "Máy bị rơi vỡ hoặc vào nước sẽ không được bảo hành.\n"
)
(root / "bao_hanh" / "quy_dinh_bao_hanh.md").write_text(md_text, encoding="utf-8")

# ----------------------------------------------------------------------
# 2. File .docx
# ----------------------------------------------------------------------
word = WordDocument()
word.add_heading("Chính sách đổi trả", level=1)
word.add_paragraph("Đổi mới trong 30 ngày đầu nếu máy lỗi do nhà sản xuất.")
word.add_paragraph("Hoàn tiền trong 7 ngày làm việc kể từ khi nhận lại máy.")
word.save(str(root / "doi_tra" / "chinh_sach_doi_tra.docx"))

# ----------------------------------------------------------------------
# 3. File .csv (có dòng tiêu đề + 3 dòng dữ liệu)
# ----------------------------------------------------------------------
csv_text = (
    "ten_may,ram_gb,pin_mah\n"
    "Xiaomi 14,12,4610\n"
    "Galaxy S24,8,4000\n"
    "iPhone 15,6,3349\n"
)
(root / "tip_su_dung" / "so_sanh_nhanh.csv").write_text(csv_text, encoding="utf-8")

# ----------------------------------------------------------------------
# 4. File .json
# ----------------------------------------------------------------------
data = {"ten_may": "Xiaomi 14", "chip": "Snapdragon 8 Gen 3", "pin_mah": 4610}
json_text = json.dumps(data, ensure_ascii=False, indent=2)
(root / "tip_su_dung" / "thong_so_mau.json").write_text(json_text, encoding="utf-8")

# ----------------------------------------------------------------------
# 5. File .html (có thẻ, để xem reader có bóc thẻ hay không)
# ----------------------------------------------------------------------
html_text = (
    "<html><body>"
    "<h1>Mẹo dùng pin</h1>"
    "<p>Không để pin cạn dưới 20 phần trăm thường xuyên.</p>"
    "</body></html>"
)
(root / "tip_su_dung" / "gioi_thieu.html").write_text(html_text, encoding="utf-8")

print("Đã tạo xong các file mẫu trong:", root)
for path_value in sorted(root.rglob("*")):
    if path_value.is_file():
        print("  ", path_value.relative_to(root).as_posix())
print("Nhớ chép thêm file .pdf vào bao_hanh/ và file .txt vào tip_su_dung/")